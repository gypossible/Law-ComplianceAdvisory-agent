"""
FastAPI 主入口
法律合规 AI 助手后端服务
"""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import asyncio
import logging
from pathlib import Path
from contextlib import asynccontextmanager

from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from typing import List
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse
from pydantic import BaseModel

from backend.llm import stream_chat, simple_chat, check_ollama_status
from backend.rag import search, add_document, get_stats, load_builtin_docs, delete_document
from backend.mindmap import generate_mindmap, create_fallback_mindmap, validate_mindmap_markdown

# 配置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).parent


@asynccontextmanager
async def lifespan(app: FastAPI):
    """应用启动时初始化知识库"""
    logger.info("🚀 启动法律合规 AI 助手...")
    
    # 检查 Ollama 状态
    status = await check_ollama_status()
    if status["status"] == "ok":
        logger.info(f"✅ Ollama 运行正常，已加载模型: {status['models']}")
    else:
        logger.warning(f"⚠️  Ollama 状态异常: {status}")
    
    # 加载内置法律文档
    logger.info("📚 正在加载法律知识库...")
    try:
        total = load_builtin_docs()
        stats = get_stats()
        logger.info(f"✅ 知识库加载完成: {stats['document_count']} 份文档, {stats['total_chunks']} 个片段")
    except Exception as e:
        logger.error(f"知识库加载失败: {e}")
    
    yield
    
    logger.info("👋 服务关闭")


app = FastAPI(
    title="法律合规 AI 助手",
    description="基于 DeepSeek R1 8B 的法律合规智能问答系统",
    version="1.0.0",
    lifespan=lifespan
)

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ======================== 数据模型 ========================

class ChatRequest(BaseModel):
    message: str
    history: list = []
    use_rag: bool = True


class MindmapRequest(BaseModel):
    question: str
    answer: str


# ======================== API 路由 ========================

@app.get("/api/health")
async def health_check():
    """健康检查"""
    ollama_status = await check_ollama_status()
    kb_stats = get_stats()
    return {
        "status": "ok",
        "ollama": ollama_status,
        "knowledge_base": kb_stats
    }


@app.post("/api/chat")
async def chat(request: ChatRequest):
    """流式对话接口"""
    
    if not request.message.strip():
        raise HTTPException(status_code=400, detail="消息不能为空")
    
    # RAG 检索
    context = ""
    if request.use_rag:
        try:
            results = search(request.message, n_results=4)
            if results:
                context_parts = []
                for r in results:
                    context_parts.append(f"【{r['source']}】\n{r['content']}")
                context = "\n\n---\n\n".join(context_parts)
        except Exception as e:
            logger.warning(f"RAG 检索失败: {e}")
    
    async def generate():
        try:
            async for chunk in stream_chat(
                user_message=request.message,
                context=context,
                history=request.history
            ):
                yield f"data: {chunk}\n\n"
        except Exception as e:
            logger.error(f"生成失败: {e}")
            yield f"data: [错误] 模型调用失败: {str(e)}\n\n"
        finally:
            yield "data: [DONE]\n\n"
    
    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no"
        }
    )


@app.post("/api/mindmap")
async def create_mindmap(request: MindmapRequest):
    """生成思维导图数据"""
    
    if not request.question or not request.answer:
        raise HTTPException(status_code=400, detail="问题和回答不能为空")
    
    try:
        markdown = await generate_mindmap(request.question, request.answer)
        
        if not validate_mindmap_markdown(markdown):
            markdown = create_fallback_mindmap(request.question, request.answer)
        
        return {"markdown": markdown, "status": "ok"}
    
    except Exception as e:
        logger.error(f"思维导图生成失败: {e}")
        # 返回备用思维导图
        fallback = create_fallback_mindmap(request.question, request.answer)
        return {"markdown": fallback, "status": "fallback"}


def parse_document_content(content_bytes: bytes, ext: str, filename: str) -> str:
    """统一文档解析函数，支持 PDF / DOCX / MD / TXT"""
    import io
    
    if ext == ".pdf":
        try:
            import PyPDF2
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content_bytes))
            pages = []
            for i, page in enumerate(pdf_reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    pages.append(f"[第{i+1}页]\n{text.strip()}")
            if not pages:
                raise ValueError("PDF 中未提取到文字内容（可能是扫描版图片 PDF）")
            return "\n\n".join(pages)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"PDF 解析失败: {e}")
    
    elif ext in (".docx", ".doc"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(content_bytes))
            sections = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    # 保留标题层级
                    if para.style.name.startswith('Heading'):
                        level = para.style.name.replace('Heading ', '')
                        prefix = '#' * int(level) if level.isdigit() else '##'
                        sections.append(f"{prefix} {text}")
                    else:
                        sections.append(text)
            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        sections.append(row_text)
            if not sections:
                raise ValueError("Word 文档内容为空")
            return "\n\n".join(sections)
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Word 解析失败: {e}")
    
    elif ext == ".md":
        # Markdown 原文保留（向量化时去除符号会自然处理）
        return content_bytes.decode("utf-8", errors="ignore")
    
    else:  # .txt
        return content_bytes.decode("utf-8", errors="ignore")


@app.post("/api/upload")
async def upload_document(
    file: UploadFile = File(...),
    doc_type: str = Form(default="custom")
):
    """上传单个文档到知识库（支持 PDF / DOCX / MD / TXT）"""
    
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名不能为空")
    
    allowed_types = [".txt", ".md", ".pdf", ".docx"]
    ext = Path(file.filename).suffix.lower()
    
    if ext not in allowed_types:
        raise HTTPException(
            status_code=400,
            detail=f"不支持的文件类型。支持格式：{' / '.join(t.upper().lstrip('.') for t in allowed_types)}"
        )
    
    try:
        content_bytes = await file.read()
        
        if len(content_bytes) > 100 * 1024 * 1024:  # 100MB 限制
            raise HTTPException(status_code=400, detail="文件大小不能超过 100MB")
        
        content = parse_document_content(content_bytes, ext, file.filename)
        
        if not content.strip():
            raise HTTPException(status_code=400, detail="文档解析后内容为空")
        
        source = Path(file.filename).stem
        chunks = add_document(content, source, doc_type)
        
        char_count = len(content)
        return {
            "status": "ok",
            "message": f"成功导入 '{file.filename}'",
            "filename": file.filename,
            "source": source,
            "chunks": chunks,
            "char_count": char_count,
            "doc_type": doc_type
        }
    
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"文档上传失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/upload/batch")
async def upload_batch(
    files: List[UploadFile] = File(...),
    doc_type: str = Form(default="custom")
):
    """批量上传多个文档"""
    if not files:
        raise HTTPException(status_code=400, detail="请选择至少一个文件")
    
    if len(files) > 10:
        raise HTTPException(status_code=400, detail="单次批量上传不超过 10 个文件")
    
    results = []
    errors = []
    
    for file in files:
        try:
            content_bytes = await file.read()
            ext = Path(file.filename).suffix.lower()
            allowed_types = [".txt", ".md", ".pdf", ".docx"]
            
            if ext not in allowed_types:
                errors.append({"filename": file.filename, "error": "不支持的文件类型"})
                continue
            
            if len(content_bytes) > 100 * 1024 * 1024:  # 100MB
                errors.append({"filename": file.filename, "error": "文件超过100MB限制"})
                continue
            
            content = parse_document_content(content_bytes, ext, file.filename)
            
            if not content.strip():
                errors.append({"filename": file.filename, "error": "内容为空"})
                continue
            
            source = Path(file.filename).stem
            chunks = add_document(content, source, doc_type)
            results.append({
                "filename": file.filename,
                "source": source,
                "chunks": chunks,
                "status": "ok"
            })
        except HTTPException as e:
            errors.append({"filename": file.filename, "error": e.detail})
        except Exception as e:
            errors.append({"filename": file.filename, "error": str(e)})
    
    return {
        "status": "ok",
        "total": len(files),
        "success": len(results),
        "failed": len(errors),
        "results": results,
        "errors": errors
    }


@app.delete("/api/knowledge-base/{source}")
async def delete_doc(source: str):
    """从知识库删除文档"""
    try:
        deleted = delete_document(source)
        if deleted:
            return {"status": "ok", "message": f"已删除文档 '{source}'"}
        else:
            raise HTTPException(status_code=404, detail=f"未找到文档 '{source}'")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/knowledge-base")
async def get_knowledge_base():
    """获取知识库信息"""
    return get_stats()


# ======================== 静态文件 ========================

frontend_dir = BASE_DIR.parent / "frontend"

if frontend_dir.exists():
    app.mount("/static", StaticFiles(directory=str(frontend_dir)), name="static")

    @app.get("/")
    async def index():
        return FileResponse(str(frontend_dir / "index.html"))


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(
        "backend.main:app",
        host="0.0.0.0",
        port=8000,
        reload=True,
        log_level="info"
    )
